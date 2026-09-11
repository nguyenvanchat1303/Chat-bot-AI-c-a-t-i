import io
import os
from datetime import datetime

import docx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt
from dotenv import load_dotenv
from google import genai
from google.genai import types
import pymupdf as fitz
import streamlit as st

# Tải biến môi trường
load_dotenv()

# Cấu hình giao diện Wide Mode
st.set_page_config(
    page_title="VKSND Khu vực 4 - Nghệ An AI", page_icon="⚖️", layout="wide"
)


# QUẢN LÝ CLIENT AN TOÀN TRÁNH LỖI CLOSED CLIENT
@st.cache_resource
def get_genai_client():
    return genai.Client(api_key=os.getenv("GOOGLE_API_KEY"))


client = get_genai_client()

# CẬP NHẬT TÊN MODEL CHUẨN TẠI ĐÂY
GEMINI_MODEL = "gemini-3.6-flash"

st.title("⚖️ Viện Kiểm sát Nhân dân Khu vực 4 - Nghệ An AI")
st.caption("Hỗ trợ tra cứu, soạn thảo Cáo trạng và bài phát biểu")
st.divider()

SYSTEM_PROMPT = """Bạn là một trợ lý pháp lý chuyên nghiệp, hỗ trợ Kiểm sát viên tại Việt Nam.
Khi soạn thảo Cáo trạng hoặc bài phát biểu, bạn PHẢI:
- Tuân thủ đúng thể thức văn bản tố tụng hình sự Việt Nam (Quốc hiệu, tiêu ngữ, phần nhận thấy, phần nhận định, phần quyết định...)
- Chỉ dựa trên thông tin, tình tiết được cung cấp, KHÔNG tự bịa thêm tình tiết hay chứng cứ không có
- Nếu thiếu thông tin cần thiết, ghi rõ [CẦN BỔ SUNG: ...] tại chỗ đó thay vì tự suy diễn
- Văn phong trang trọng, chính xác, đúng thuật ngữ pháp lý
- Cuối mỗi văn bản, LUÔN ghi rõ: "Đây là bản dự thảo do AI hỗ trợ soạn, cần được Kiểm sát viên có thẩm quyền kiểm tra, đối chiếu hồ sơ và phê duyệt trước khi ban hành chính thức." """


# ====== HÀM TẠO FILE WORD THEO NĐ 30 ======
def tao_file_word(
    noi_dung_van_ban, tieu_de, so_hieu="...../CT-VKS", dia_danh="Nghệ An"
):
    doc = Document()

    # Lề trang theo Nghị định 30
    section = doc.sections[0]
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)

    # BẢNG 2 CỘT: Cơ quan ban hành (trái) | Quốc hiệu tiêu ngữ (phải)
    table = doc.add_table(rows=1, cols=2)
    table.autofit = True

    # Cột trái: Cơ quan ban hành
    cell_trai = table.cell(0, 0)
    p_capttren = cell_trai.paragraphs[0]
    p_capttren.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_capttren = p_capttren.add_run(
        "VIỆN KIỂM SÁT NHÂN DÂN\nTỈNH NGHỆ AN"
    )
    run_capttren.font.size = Pt(13)
    run_capttren.font.name = "Times New Roman"

    p_capduoi = cell_trai.add_paragraph()
    p_capduoi.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_capduoi = p_capduoi.add_run("VIỆN KIỂM SÁT NHÂN DÂN\nKHU VỰC 4")
    run_capduoi.bold = True
    run_capduoi.font.size = Pt(13)
    run_capduoi.font.name = "Times New Roman"

    p_gach = cell_trai.add_paragraph()
    p_gach.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_gach = p_gach.add_run("________")
    run_gach.font.size = Pt(13)

    p_sohieu = cell_trai.add_paragraph()
    p_sohieu.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sohieu = p_sohieu.add_run(f"Số: {so_hieu}")
    run_sohieu.font.size = Pt(13)
    run_sohieu.font.name = "Times New Roman"

    # Cột phải: Quốc hiệu tiêu ngữ
    cell_phai = table.cell(0, 1)
    p_qh = cell_phai.paragraphs[0]
    p_qh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_qh = p_qh.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    run_qh.bold = True
    run_qh.font.size = Pt(13)
    run_qh.font.name = "Times New Roman"

    p_tn = cell_phai.add_paragraph()
    p_tn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_tn = p_tn.add_run("Độc lập - Tự do - Hạnh phúc")
    run_tn.bold = True
    run_tn.font.size = Pt(14)
    run_tn.font.name = "Times New Roman"

    p_gach2 = cell_phai.add_paragraph()
    p_gach2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_gach2 = p_gach2.add_run("---------------")
    run_gach2.font.size = Pt(13)

    p_ngay = cell_phai.add_paragraph()
    p_ngay.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hom_nay = datetime.now()
    run_ngay = p_ngay.add_run(
        f"{dia_danh}, ngày {hom_nay.day} tháng {hom_nay.month} năm {hom_nay.year}"
    )
    run_ngay.italic = True
    run_ngay.font.size = Pt(13)
    run_ngay.font.name = "Times New Roman"

    # Xóa viền bảng 1
    for cell in table._cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = tcPr.makeelement(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcBorders",
            {},
        )
        for edge in ("top", "left", "bottom", "right"):
            tag = f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{edge}"
            el = tcPr.makeelement(
                tag,
                {
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val": "nil"
                },
            )
            tcBorders.append(el)
        tcPr.append(tcBorders)

    doc.add_paragraph()

    # TÊN VĂN BẢN
    p_ten = doc.add_paragraph()
    p_ten.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_ten = p_ten.add_run(tieu_de.upper())
    run_ten.bold = True
    run_ten.font.size = Pt(16)
    run_ten.font.name = "Times New Roman"

    doc.add_paragraph()

    # NỘI DUNG CHÍNH
    for doan in noi_dung_van_ban.split("\n"):
        if doan.strip():
            p = doc.add_paragraph(doan)
            p.paragraph_format.first_line_indent = Mm(10)
            p.paragraph_format.line_spacing = 1.2
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(14)

    doc.add_paragraph()

    # NƠI NHẬN + CHỮ KÝ
    table2 = doc.add_table(rows=1, cols=2)

    cell_noinhap = table2.cell(0, 0)
    p_nn = cell_noinhap.paragraphs[0]
    run_nn = p_nn.add_run("Nơi nhận:\n- Như trên;\n- Lưu: VT, HSKS.")
    run_nn.italic = True
    run_nn.font.size = Pt(12)
    run_nn.font.name = "Times New Roman"

    cell_kyten = table2.cell(0, 1)
    p_chucvu = cell_kyten.paragraphs[0]
    p_chucvu.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_chucvu = p_chucvu.add_run("VIỆN TRƯỞNG")
    run_chucvu.bold = True
    run_chucvu.font.size = Pt(14)
    run_chucvu.font.name = "Times New Roman"

    for _ in range(3):
        cell_kyten.add_paragraph()

    p_hoten = cell_kyten.add_paragraph()
    p_hoten.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_hoten = p_hoten.add_run("[Họ tên]")
    run_hoten.bold = True
    run_hoten.font.size = Pt(14)
    run_hoten.font.name = "Times New Roman"

    # Xóa viền bảng 2
    for cell in table2._cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = tcPr.makeelement(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcBorders",
            {},
        )
        for edge in ("top", "left", "bottom", "right"):
            tag = f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{edge}"
            el = tcPr.makeelement(
                tag,
                {
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val": "nil"
                },
            )
            tcBorders.append(el)
        tcPr.append(tcBorders)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ====== SIDEBAR MENU ======
with st.sidebar:
    st.header("📄 Tải tài liệu")
    uploaded_file = st.file_uploader(
        "Tải lên file PDF hồ sơ, văn bản...", type=["pdf"], key="pdf_uploader"
    )

    if uploaded_file is not None:
        pdf_bytes = uploaded_file.read()
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        pdf_text = ""
        for page in doc:
            pdf_text += page.get_text() + "\n"
        doc.close()
        st.session_state.pdf_content = pdf_text
        st.success(
            f"Đã đọc xong file: {uploaded_file.name} ({len(pdf_text)} ký tự)"
        )

    st.markdown("---")
    st.header("🎯 Chế độ làm việc")
    che_do = st.radio(
        "Chọn việc bạn muốn làm:",
        [
            "💬 Trò chuyện tự do",
            "📋 Soạn dự thảo Cáo trạng",
            "🎤 Soạn bài phát biểu",
        ],
    )

    st.markdown("---")
    if st.button("🗑️ Xóa lịch sử trò chuyện", use_container_width=True):
        st.session_state.messages = []
        if "pdf_content" in st.session_state:
            del st.session_state.pdf_content
        st.session_state.chat = get_genai_client().chats.create(
            model=GEMINI_MODEL,
            config=types.GenerateContentConfig(
                system_instruction=SYSTEM_PROMPT
            ),
        )
        st.rerun()

# KHỞI TẠO STATE
if "chat" not in st.session_state:
    st.session_state.chat = get_genai_client().chats.create(
        model=GEMINI_MODEL,
        config=types.GenerateContentConfig(system_instruction=SYSTEM_PROMPT),
    )
    st.session_state.messages = []

pdf_note = ""
if "pdf_content" in st.session_state:
    pdf_note = f"\n\n[Nội dung tài liệu PDF đã tải lên]:\n{st.session_state.pdf_content}"


# ====== 1. CHẾ ĐỘ: SOẠN DỰ THẢO CÁO TRẠNG ======
if che_do == "📋 Soạn dự thảo Cáo trạng":
    col1, col2 = st.columns([1, 1], gap="large")

    with col1:
        st.subheader("📋 Nhập thông tin chi tiết Cáo trạng")
        with st.form("form_cao_trang"):
            # 1. Căn cứ tố tụng
            st.markdown("**1. PHẦN CĂN CỨ TỐ TỤNG**")
            so_kt_vu_an = st.text_input(
                "Quyết định khởi tố vụ án số, ngày, cơ quan ban hành:",
                placeholder="VD: Số 12/QĐ-CSĐT ngày 10/01/2026 của Cơ quan CSĐT...",
            )
            so_kt_bi_can = st.text_input(
                "Quyết định khởi tố bị can số, ngày, cơ quan ban hành:",
                placeholder="VD: Số 15/QĐ-CSĐT ngày 12/01/2026...",
            )
            so_kldt = st.text_input(
                "Kết luận điều tra số, ngày, cơ quan ban hành:",
                placeholder="VD: Số 05/KLĐT ngày 01/03/2026...",
            )

            # 2. Nội dung vụ án
            st.markdown("**2. PHẦN NỘI DUNG VỤ ÁN**")
            dien_bien = st.text_area(
                "Diễn biến hành vi phạm tội (theo trình tự thời gian):",
                height=120,
            )
            vat_chung = st.text_input(
                "Thu giữ, tạm giữ đồ vật, tài liệu, vật chứng (nếu có):",
                placeholder="VD: Thu giữ 01 xe máy Honda Wave BKS...",
            )
            dan_su = st.text_input(
                "Trách nhiệm dân sự, bồi thường thiệt hại (nếu có):",
                placeholder="VD: Bị hại yêu cầu bồi thường 50.000.000đ...",
            )

            # 3. Thông tin bị can & Kết luận
            st.markdown("**3. KẾT LUẬN & THÔNG TIN BỊ CAN**")
            ten_bi_can = st.text_input("Họ tên bị can:")
            thong_tin_lai_lich = st.text_input(
                "Giới tính, ngày/năm sinh, nơi cư trú, tiền án tiền sự, biện pháp ngăn chặn:",
                placeholder="VD: Nam, sinh 1990, không tiền án tiền sự, đang áp dụng Tạm giam...",
            )
            toi_danh = st.text_input(
                "Tội danh đề nghị truy tố:", placeholder="VD: Trộm cắp tài sản"
            )
            dieu_luat = st.text_input(
                "Điều, khoản BLHS áp dụng:",
                placeholder="VD: Khoản 2 Điều 173 Bộ luật Hình sự",
            )
            tang_nhe_giam_nhe = st.text_input(
                "Tình tiết tăng nặng, giảm nhẹ TNHS:",
                placeholder="VD: Giảm nhẹ theo điểm s khoản 1 Điều 51 BLHS...",
            )

            # 4. Quyết định
            st.markdown("**4. PHẦN QUYẾT ĐỊNH**")
            toa_an_truy_to = st.text_input(
                "Tòa án nhân dân đề nghị truy tố xét xử:",
                placeholder="VD: Tòa án nhân dân khu vực 4 - Nghệ An",
            )
            so_ho_so = st.text_input(
                "Số lượng tập, tờ hồ sơ kèm theo:",
                placeholder="VD: 01 tập, 150 tờ",
            )

            submitted = st.form_submit_button(
                "🚀 Soạn dự thảo Cáo trạng",
                type="primary",
                use_container_width=True,
            )

    with col2:
        st.subheader("📄 Kết quả dự thảo")
        if submitted:
            prompt = f"""Hãy soạn một bản dự thảo Cáo trạng dựa trên thông tin được cung cấp bên dưới, tuân thủ CHÍNH XÁC cấu trúc và nội dung 8 phần sau:

1. PHẦN CĂN CỨ (mở đầu):
- Căn cứ các Điều 41, 236, 239 và 243 của Bộ luật Tố tụng hình sự;
- Căn cứ Quyết định khởi tố vụ án hình sự: {so_kt_vu_an or '[CẦN BỔ SUNG: Số, ngày QĐ khởi tố vụ án]'}
- Căn cứ Quyết định khởi tố bị can: {so_kt_bi_can or '[CẦN BỔ SUNG: Số, ngày QĐ khởi tố bị can]'}
- Căn cứ bản Kết luận điều tra vụ án hình sự đề nghị truy tố: {so_kldt or '[CẦN BỔ SUNG: Số, ngày Kết luận điều tra]'}

2. "Trên cơ sở kết quả điều tra đã xác định được như sau:"

3. PHẦN NỘI DUNG:
- Diễn biến hành vi phạm tội: {dien_bien or '[CẦN BỔ SUNG: Diễn biến hành vi phạm tội]'}
- Phân tích, đánh giá tình tiết tăng nặng, giảm nhẹ và tình tiết khác có ý nghĩa đối với vụ án.
- Việc thu giữ, tạm giữ đồ vật, tài liệu, xử lý vật chứng: {vat_chung or 'Không có'}
- Phần dân sự: {dan_su or 'Không có phần dân sự trong vụ án'}

4. "Căn cứ vào các tình tiết và chứng cứ nêu trên,"

5. KẾT LUẬN:
- Tổng hợp hành vi phạm tội: nêu ngắn gọn hành vi, tính chất, mức độ, hậu quả, vai trò bị can.
- Câu bắt buộc: "Như vậy có đủ căn cứ để xác định bị can có lý lịch dưới đây đã phạm tội {toi_danh or '[CẦN BỔ SUNG: Tên tội danh]'} như sau:"
- Thông tin bị can:
  + Họ và tên: {ten_bi_can or '[CẦN BỔ SUNG: Họ tên bị can]'}
  + Lý lịch, tiền án tiền sự, biện pháp ngăn chặn: {thong_tin_lai_lich or '[CẦN BỔ SUNG: Lý lịch bị can]'}
- Khẳng định: Bị can {ten_bi_can} đã phạm tội {toi_danh}, quy định tại {dieu_luat or '[CẦN BỔ SUNG: Điều khoản BLHS]'}
- Tình tiết tăng nặng, giảm nhẹ trách nhiệm hình sự: {tang_nhe_giam_nhe or '[CẦN BỔ SUNG: Điểm, khoản áp dụng]'}

6. "Bởi các lẽ trên,"

7. PHẦN "QUYẾT ĐỊNH" (viết in hoa):
1. Truy tố ra trước {toa_an_truy_to or 'Tòa án nhân dân khu vực... [CẦN BỔ SUNG]'} để xét xử bị can {ten_bi_can} về tội {toi_danh} theo quy định tại {dieu_luat}.
2. Kèm theo Cáo trạng có: Hồ sơ vụ án gồm {so_ho_so or '...tập, ...tờ [CẦN BỔ SUNG]'}. Bản kê vật chứng (nếu có). Danh sách những người Viện kiểm sát đề nghị Tòa án triệu tập đến phiên tòa.

8. PHẦN CUỐI: KHÔNG tự tạo phần Nơi nhận và Chữ ký (để trống hoàn toàn).

{pdf_note}
LƯU Ý QUAN TRỌNG: Với các thông tin chưa có trong dữ liệu nhập, PHẢI ghi rõ [CẦN BỔ SUNG: ...] tại đúng vị trí đó, tuyệt đối không tự bịa ra số liệu hay ngày tháng."""

            with st.spinner("Đang phân tích dữ liệu và soạn Cáo trạng..."):
                response = st.session_state.chat.send_message(prompt)
                st.write(response.text)

                file_word = tao_file_word(response.text, "CÁO TRẠNG")
                st.download_button(
                    "💾 Tải về file Word (.docx)",
                    file_word,
                    file_name="du_thao_cao_trang.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True,
                )


# ====== 2. CHẾ ĐỘ: SOẠN BÀI PHÁT BIỂU ======
elif che_do == "🎤 Soạn bài phát biểu":
    col1, col2 = st.columns([1, 1], gap="large")

    with col1:
        st.subheader("🎤 Nhập thông tin bài phát biểu")
        with st.form("form_phat_bieu"):
            loai_phat_bieu = st.selectbox(
                "Loại bài phát biểu",
                [
                    "Luận tội tại phiên tòa",
                    "Phát biểu tranh luận",
                    "Phát biểu khác",
                ],
            )
            vu_viec = st.text_area(
                "Tóm tắt vụ việc / nội dung cần phát biểu", height=150
            )
            diem_nhan_manh = st.text_area(
                "Các điểm cần nhấn mạnh (nếu có)", height=100
            )
            do_dai = st.select_slider(
                "Độ dài mong muốn",
                options=["Ngắn gọn", "Vừa phải", "Chi tiết"],
                value="Vừa phải",
            )
            submitted2 = st.form_submit_button(
                "🚀 Soạn bài phát biểu",
                type="primary",
                use_container_width=True,
            )

    with col2:
        st.subheader("📄 Kết quả dự thảo")
        if submitted2:
            prompt = f"""Hãy soạn một bài phát biểu loại "{loai_phat_bieu}" với thông tin sau:
- Tóm tắt vụ việc: {vu_viec}
- Các điểm cần nhấn mạnh: {diem_nhan_manh}
- Độ dài mong muốn: {do_dai}
{pdf_note}

Soạn với văn phong trang trọng, mạch lạc, phù hợp để trình bày tại phiên tòa hoặc cuộc họp chính thức."""

            with st.spinner("Đang soạn bài phát biểu..."):
                response = st.session_state.chat.send_message(prompt)
                st.write(response.text)

                file_word = tao_file_word(response.text, "BÀI PHÁT BIỂU")
                st.download_button(
                    "💾 Tải về file Word (.docx)",
                    file_word,
                    file_name="du_thao_bai_phat_bieu.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True,
                )


# ====== 3. CHẾ ĐỘ: TRÒ CHUYỆN TỰ DO ======
else:
    st.subheader("💬 Hỏi đáp & Trợ lý pháp lý")
    for msg in st.session_state.messages:
        avatar = "🧑" if msg["role"] == "user" else "⚖️"
        with st.chat_message(msg["role"], avatar=avatar):
            st.write(msg["content"])

    user_input = st.chat_input(
        "Đặt câu hỏi pháp lý hoặc yêu cầu tra cứu...", key="main_chat_input"
    )

    if user_input:
        final_input = (
            user_input + pdf_note
            if "pdf_content" in st.session_state
            else user_input
        )

        st.session_state.messages.append(
            {"role": "user", "content": user_input}
        )
        with st.chat_message("user", avatar="🧑"):
            st.write(user_input)

        with st.chat_message("assistant", avatar="⚖️"):
            with st.spinner("Đang tra cứu..."):
                response = st.session_state.chat.send_message(final_input)
                st.write(response.text)
        st.session_state.messages.append(
            {"role": "assistant", "content": response.text}
        )